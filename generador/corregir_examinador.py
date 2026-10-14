"""
Corrige en los formularios del OEC el examinador mal copiado de otro candidato
(lista en informes/examinador_mal_copiado.csv):

    C08: fila "EXAMINADOR" (nombre, cédula, celular)
    C11: nombre y C.C. del examinador

Solo cambia esas celdas, editando el XML del .xlsx (el resto del archivo
queda igual). Antes guarda el original en respaldos/correcciones/.
Al final vuelve a leer cada archivo para comprobar el cambio.

USO:
    python corregir_examinador.py            # aplica las correcciones
    python corregir_examinador.py --prueba   # solo muestra lo que haría
"""

import csv
import os
import re
import shutil
import sys
import warnings
import zipfile
from xml.sax.saxutils import escape

from openpyxl import load_workbook

import ajustes
import base
import base_escritura
import config

warnings.filterwarnings("ignore")
LISTA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "informes", "examinador_mal_copiado.csv")
RESPALDOS = os.path.join(base_escritura.CARPETA_RESPALDOS, "correcciones")


def poner_celdas(ruta, hoja, cambios):
    """cambios: {"B18": "texto", ...} en la hoja indicada. Conserva el estilo de cada celda."""
    with zipfile.ZipFile(ruta) as zin:
        hoja_xml = base_escritura.ruta_hoja(zin, hoja)
        entradas = [(i, zin.read(i.filename)) for i in zin.infolist()]
    xml = dict((i.filename, d) for i, d in entradas)[hoja_xml].decode("utf-8")
    for ref, valor in cambios.items():
        nueva_t = f'<is><t xml:space="preserve">{escape(valor)}</t></is>'
        patron = re.compile(rf'<c\b([^>]*\br="{ref}"[^>]*?)(?:/>|>.*?</c>)', re.S)
        m = patron.search(xml)
        if m:
            atributos = re.sub(r'\s+t="[^"]*"', "", m.group(1))
            xml = xml[:m.start()] + f'<c{atributos} t="inlineStr">{nueva_t}</c>' + xml[m.end():]
        else:
            fila = re.search(r"\d+", ref).group(0)
            f = re.search(rf'<row\b[^>]*\br="{fila}"[^>]*?(?:/>|>)', xml)
            if not f:
                raise ValueError(f"no existe la fila {fila}")
            apertura = f.group(0)
            celda = f'<c r="{ref}" t="inlineStr">{nueva_t}</c>'
            if apertura.endswith("/>"):
                xml = xml.replace(apertura, apertura[:-2] + ">" + celda + "</row>", 1)
            else:
                xml = xml.replace(apertura, apertura + celda, 1)
    tmp = ruta + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, datos in entradas:
            zout.writestr(info, xml.encode("utf-8") if info.filename == hoja_xml else datos)
    os.replace(tmp, ruta)


def ubicar_c08(ws):
    for fila in ws.iter_rows(min_row=10, max_row=30):
        for c in fila:
            if base.normalizar(c.value) == "EXAMINADOR":
                n = c.row
                return {"nombre": f"B{n}", "cedula": f"E{n}", "celular": f"G{n}"}
    raise ValueError("no se encontró la fila EXAMINADOR")


def ubicar_c11(ws):
    nombre = cedula = None
    for fila in ws.iter_rows(min_row=1, max_row=15):
        for c in fila:
            v = base.normalizar(c.value)
            if v.startswith("APELLIDOS Y NOMBRES DEL EXAMINADOR"):
                nombre = f"H{c.row}"
            if v in ("C.C.", "C.C", "CC", "C.C.:"):
                cedula = f"B{c.row}"
    if not (nombre and cedula):
        raise ValueError("no se encontraron las celdas del examinador")
    return {"nombre": nombre, "cedula": cedula}


def main():
    prueba = "--prueba" in sys.argv
    aj = ajustes.leer()
    filas = list(csv.DictReader(open(LISTA, encoding="utf-8-sig"), delimiter=";"))
    hechos = errores = 0
    for f in filas:
        perfil, codigo, formulario = f["perfil"], f["codigo"], f["formulario"]
        exa = ajustes.examinador(aj, aj["perfiles"][perfil]["examinador"])
        carpeta = config.carpeta_formularios(perfil, codigo)
        prefijo = "C0?8" if formulario == "C08" else "C11"
        archivo = next(a for a in os.listdir(carpeta) if re.match(prefijo + r"[_\-]", a, re.I) and a.endswith(".xlsx"))
        ruta = os.path.join(carpeta, archivo)
        try:
            wb = load_workbook(ruta, read_only=False, data_only=False)
            ws = wb[wb.sheetnames[0]] if formulario == "C11" else (wb["HOJA DE TRABAJO"] if "HOJA DE TRABAJO" in wb.sheetnames else wb.worksheets[0])
            celdas = ubicar_c08(ws) if formulario == "C08" else ubicar_c11(ws)
            antes = {k: ws[v].value for k, v in celdas.items()}
            cambios = {celdas["nombre"]: exa["nombre"], celdas["cedula"]: exa["cedula"]}
            if "celular" in celdas:
                cambios[celdas["celular"]] = exa["celular"]
            print(f"{'(prueba) ' if prueba else ''}{codigo} {formulario}: {antes} -> {exa['nombre']} / {exa['cedula']}")
            if prueba:
                continue
            destino = os.path.join(RESPALDOS, perfil, codigo)
            os.makedirs(destino, exist_ok=True)
            if not os.path.exists(os.path.join(destino, archivo)):
                shutil.copy2(ruta, os.path.join(destino, archivo))
            poner_celdas(ruta, ws.title, cambios)
            ws2 = load_workbook(ruta, data_only=True)[ws.title]
            if str(ws2[celdas["cedula"]].value) != exa["cedula"] or base.normalizar(ws2[celdas["nombre"]].value) != base.normalizar(exa["nombre"]):
                raise RuntimeError("al volver a leer no quedó el examinador correcto")
            hechos += 1
        except Exception as e:
            errores += 1
            print(f"   ✗ {codigo} {formulario}: {e}")
    print(f"\nCorregidos: {hechos}. Con error: {errores}. Respaldos en {RESPALDOS}")


if __name__ == "__main__":
    main()
